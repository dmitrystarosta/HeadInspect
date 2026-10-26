"""Excel report export (backend/app/report.py + the report.xlsx route).

Builds workbooks from synthetic jobs and loads the produced bytes back with
openpyxl to assert real structure (sheets, headers, rows), then exercises the
route for its status codes and download headers. No network, no real audit.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook
from fastapi.testclient import TestClient

from app.jobs import Job, job_manager
from app.main import app
from app.models import (
    CanonicalData,
    ImageItem,
    ImagesData,
    MetaData,
    OpenGraphData,
    PageResult,
    SchemaData,
)
from app.report import build_report_workbook, report_filename


EXPECTED_SHEETS = ["Сводка", "Open Graph", "Meta", "Canonical", "Schema", "Images", "Sitemap", "Проблемы"]


def _page(url, **kw) -> PageResult:
    kw.setdefault("requested_url", url)
    kw.setdefault("status_code", 200)
    return PageResult(url=url, **kw)


def _completed_job() -> Job:
    job = Job(job_id="done1", requested_url="https://example.ru/")
    job.normalized_url = "https://example.ru/"
    job.status = "completed"
    job.discovered_urls = 2
    job.checked_urls = 2
    job.results = [
        _page(
            "https://example.ru/",
            open_graph=OpenGraphData(title="Главная", image="https://example.ru/og.jpg",
                                     image_accessible=True),
            meta=MetaData(title="Главная", description="Описание", warnings=["Нет meta keywords"]),
            canonical=CanonicalData(resolved_url="https://example.ru/", source="html", is_self=True),
            schema_data=SchemaData(json_ld_count=1, types=["WebSite"]),
            errors=["Нет og:description"],
        ),
        _page(
            "https://example.ru/about",
            meta=MetaData(title="О нас", description="Про нас"),
            canonical=CanonicalData(resolved_url="https://example.ru/about", source="html", is_self=True),
            schema_data=SchemaData(json_ld_count=0),
        ),
    ]
    return job


def _load(job: Job):
    return load_workbook(BytesIO(build_report_workbook(job)))


def test_workbook_has_all_sheets_in_order():
    wb = _load(_completed_job())
    assert wb.sheetnames == EXPECTED_SHEETS


def test_module_sheet_headers_and_one_row_per_page():
    wb = _load(_completed_job())
    meta = wb["Meta"]
    header = [c.value for c in meta[1]]
    assert header[:3] == ["URL", "HTTP", "title"]
    assert "Ошибки" in header and "Предупреждения" in header
    # Header row + two content pages.
    assert meta.max_row == 3


def test_problems_sheet_lists_errors_and_warnings_flat():
    wb = _load(_completed_job())
    rows = [tuple(r) for r in wb["Проблемы"].iter_rows(values_only=True)]
    assert rows[0] == ("URL", "Модуль", "Тип", "Сообщение")
    body = rows[1:]
    assert ("https://example.ru/", "Open Graph", "Ошибка", "Нет og:description") in body
    assert ("https://example.ru/", "Meta", "Предупреждение", "Нет meta keywords") in body


def test_summary_counts_and_partial_reason():
    job = _completed_job()
    job.status = "completed_partial"
    job.partial_reason = "Достигнут лимит времени."
    wb = _load(job)
    pairs = {r[0]: r[1] for r in wb["Сводка"].iter_rows(values_only=True) if r[0]}
    assert pairs["Статус"] == "Завершён частично"
    assert pairs["Страниц проверено"] == "2"
    assert pairs["Проверка завершена не полностью"] == "Достигнут лимит времени."


def test_redirect_duplicate_is_deduplicated_in_content_sheet():
    # A page reached directly AND the same final URL reached via a redirect
    # must collapse to a single content row (mirrors uniqueContentPages).
    job = _completed_job()
    job.results.append(
        _page("https://example.ru/", requested_url="https://example.ru/index.html")
    )
    wb = _load(job)
    # Still exactly two unique content pages ( / and /about ).
    assert wb["Meta"].max_row == 3
    # Sitemap keeps every discovered URL, including the redirect entry.
    assert wb["Sitemap"].max_row == 4


def test_check_failed_page_is_availability_problem_not_content_row():
    job = _completed_job()
    job.results.append(
        PageResult(url="https://example.ru/dead", requested_url="https://example.ru/dead",
                   status_code=None, check_failed=True, check_reason="network",
                   check_error="Не удалось получить страницу")
    )
    wb = _load(job)
    # Not added to content module sheets (still 2 content rows + header).
    assert wb["Meta"].max_row == 3
    problems = [tuple(r) for r in wb["Проблемы"].iter_rows(values_only=True)]
    assert ("https://example.ru/dead", "Доступность", "Не проверено",
            "Не удалось получить страницу") in problems


def test_zero_results_produces_valid_workbook():
    job = Job(job_id="empty1", requested_url="https://example.ru/")
    job.normalized_url = "https://example.ru/"
    job.status = "completed"
    wb = _load(job)
    assert wb.sheetnames == EXPECTED_SHEETS
    # Every list sheet is header-only.
    for name in ["Open Graph", "Meta", "Canonical", "Schema", "Images", "Sitemap", "Проблемы"]:
        assert wb[name].max_row == 1


def test_report_filename_uses_host_and_date():
    name = report_filename(_completed_job())
    assert name.startswith("headinspect-example.ru-")
    assert name.endswith(".xlsx")


def test_formula_injection_is_neutralized():
    # Site-controlled strings that look like formulas must be written as text,
    # never as executable formula cells.
    job = _completed_job()
    job.results = [
        _page(
            "https://example.ru/evil",
            open_graph=OpenGraphData(title="=HYPERLINK(\"http://evil\")", image=None),
            meta=MetaData(title="+1+1", description="@SUM(A1)", warnings=["=cmd|' /C calc'!A0"]),
            canonical=CanonicalData(resolved_url="https://example.ru/evil", source="html"),
            schema_data=SchemaData(json_ld_count=0),
            errors=["-dangerous"],
        ),
    ]
    wb = _load(job)
    # No cell anywhere is a formula.
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", f"formula leaked into {ws.title}!{cell.coordinate}"
    # The risky title is present, as literal text (leading apostrophe marker).
    og = wb["Open Graph"]
    titles = [r[2] for r in og.iter_rows(min_row=2, values_only=True)]
    assert any(isinstance(t, str) and t.startswith("'=HYPERLINK") for t in titles)


def test_open_graph_sheet_and_problems_exclude_http_availability_lines():
    # A 404 that still returned HTML: top-level errors carry both "HTTP 404"
    # (availability) and a real OG finding. The OG sheet/rows must show only
    # the OG finding; the HTTP line belongs to «Доступность».
    job = _completed_job()
    job.results = [
        _page(
            "https://example.ru/notfound",
            status_code=404,
            open_graph=OpenGraphData(title=None, image=None),
            meta=MetaData(title="X"),
            canonical=CanonicalData(source="none"),
            schema_data=SchemaData(json_ld_count=0),
            errors=["HTTP 404", "Нет og:title"],
        ),
    ]
    wb = _load(job)
    og_row = list(wb["Open Graph"].iter_rows(min_row=2, values_only=True))[0]
    og_errors_cell = og_row[8]
    assert "Нет og:title" in og_errors_cell
    assert "HTTP 404" not in og_errors_cell

    problems = [tuple(r) for r in wb["Проблемы"].iter_rows(values_only=True)]
    assert ("https://example.ru/notfound", "Open Graph", "Ошибка", "Нет og:title") in problems
    assert ("https://example.ru/notfound", "Доступность", "Ошибка", "HTTP 404") in problems
    # The HTTP line is never attributed to Open Graph.
    assert not any(p[1] == "Open Graph" and p[3] == "HTTP 404" for p in problems)


def test_images_sheet_and_problems_rows():
    job = _completed_job()
    job.results[0].images = ImagesData(
        items=[
            ImageItem(src="https://example.ru/ok.png", source="img", has_alt=True,
                      alt="Логотип", width_attr="100", height_attr="40",
                      format_guess="PNG", checked=True, accessible=True,
                      status_code=200, bytes=2048, reuse_count=3),
            ImageItem(src="https://example.ru/gone.png", source="img", has_alt=False,
                      checked=True, accessible=False, status_code=404,
                      errors=["HTTP 404"], warnings=["нет атрибута alt"]),
        ],
        total=2, checked_count=2,
        errors=["https://example.ru/gone.png: HTTP 404"],
        warnings=["https://example.ru/gone.png: нет атрибута alt"],
    )
    wb = _load(job)
    img = wb["Images"]
    header = [c.value for c in img[1]]
    assert header[:4] == ["Страница", "URL изображения", "Источник", "alt"]
    # Two image rows + header.
    assert img.max_row == 3
    problems = [tuple(r) for r in wb["Проблемы"].iter_rows(values_only=True)]
    assert any(p[1] == "Images" and "HTTP 404" in p[3] for p in problems)
    assert any(p[1] == "Images" and "alt" in p[3] for p in problems)


def test_sitemap_discovery_issues_appear_in_problems():
    job = _completed_job()
    job.sitemap_issues = ["Sitemap указан в robots.txt, но не был получен"]
    wb = _load(job)
    problems = [tuple(r) for r in wb["Проблемы"].iter_rows(values_only=True)]
    assert ("https://example.ru/", "Sitemap", "Предупреждение",
            "Sitemap указан в robots.txt, но не был получен") in problems


# --- route ------------------------------------------------------------------

def _client() -> TestClient:
    return TestClient(app)


def test_route_returns_xlsx_with_download_headers():
    job = _completed_job()
    job_manager.jobs[job.job_id] = job
    try:
        resp = _client().get(f"/api/audits/{job.job_id}/report.xlsx")
        assert resp.status_code == 200
        assert resp.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "attachment" in resp.headers["content-disposition"]
        # The bytes really are a loadable workbook.
        wb = load_workbook(BytesIO(resp.content))
        assert wb.sheetnames == EXPECTED_SHEETS
    finally:
        job_manager.jobs.pop(job.job_id, None)


def test_route_409_while_unfinished():
    job = Job(job_id="run1", requested_url="https://example.ru/")
    job.status = "running"
    job_manager.jobs[job.job_id] = job
    try:
        resp = _client().get(f"/api/audits/{job.job_id}/report.xlsx")
        assert resp.status_code == 409
    finally:
        job_manager.jobs.pop(job.job_id, None)


def test_route_404_for_unknown_job():
    resp = _client().get("/api/audits/does-not-exist/report.xlsx")
    assert resp.status_code == 404
